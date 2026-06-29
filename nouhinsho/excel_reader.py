from __future__ import annotations

from pathlib import Path
from typing import Iterable

import openpyxl

from .models import Order
from .rules import clean_text, format_postal_code


class ExcelReadError(ValueError):
    pass


REQUIRED_TIKTOK_COLUMNS = {
    "注文ID",
    "商品名",
    "数量",
    "受取人",
    "郵便番号",
    "都道府県",
    "市区町村",
    "町名",
    "詳細住所1",
    "電話番号",
}


def _row_value(row: tuple, index: dict[str, int], key: str) -> object:
    pos = index.get(key)
    return row[pos] if pos is not None and pos < len(row) else ""


def _address(row: tuple, index: dict[str, int]) -> str:
    pref = clean_text(_row_value(row, index, "都道府県"))
    city = clean_text(_row_value(row, index, "市区町村"))
    town = clean_text(_row_value(row, index, "町名"))
    line1 = clean_text(_row_value(row, index, "詳細住所1"))
    line2 = clean_text(_row_value(row, index, "詳細住所2"))
    # Some platforms put the whole address in 詳細住所1. Avoid repeating pieces.
    pieces = [pref, city, town]
    prefix = "".join(pieces)
    base = line1 if any(part and part in line1 for part in pieces) else prefix + line1
    if line2 and line2.replace(" ", "") not in base.replace(" ", ""):
        base = f"{base}, {line2}"
    return base


def _find_header_row(rows: list[tuple]) -> tuple[int, dict[str, int]]:
    for row_idx, row in enumerate(rows[:10]):
        values = [clean_text(x) for x in row]
        if "注文ID" in values and "商品名" in values:
            return row_idx, {name: idx for idx, name in enumerate(values) if name}
    raise ExcelReadError("Excel内に『注文ID』『商品名』のヘッダー行が見つかりません。")


def read_orders_from_excel(path: str | Path, *, source: str = "tiktok") -> list[Order]:
    workbook = openpyxl.load_workbook(path, data_only=True)
    results: list[Order] = []

    for sheet_name in workbook.sheetnames:
        ws = workbook[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        try:
            header_row, index = _find_header_row(rows)
        except ExcelReadError:
            continue

        missing = REQUIRED_TIKTOK_COLUMNS - set(index)
        if missing:
            raise ExcelReadError(f"Excel缺少字段: {', '.join(sorted(missing))}")

        for row in rows[header_row + 1 :]:
            order_no = clean_text(_row_value(row, index, "注文ID"))
            # TikTok exports often have a second explanatory row below the header.
            if not order_no or "TikTok Shop" in order_no or "固有" in order_no:
                continue
            quantity_text = clean_text(_row_value(row, index, "数量")) or "1"
            try:
                quantity = int(float(quantity_text))
            except ValueError:
                quantity = 1
            order_datetime = clean_text(_row_value(row, index, "注文の支払い日時"))
            if not order_datetime:
                order_datetime = clean_text(_row_value(row, index, "注文作成時刻"))

            results.append(
                Order(
                    source=source,
                    order_no=order_no,
                    order_datetime=order_datetime,
                    recipient=clean_text(_row_value(row, index, "受取人")),
                    postal_code=format_postal_code(_row_value(row, index, "郵便番号")),
                    address=_address(row, index),
                    phone=clean_text(_row_value(row, index, "電話番号")),
                    product_name=clean_text(_row_value(row, index, "商品名")),
                    care=clean_text(_row_value(row, index, "バリエーション")),
                    sku=clean_text(_row_value(row, index, "セラーSKU")),
                    quantity=quantity,
                )
            )
    if not results:
        raise ExcelReadError("可出力の注文が見つかりませんでした。")
    return results
